import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette - Modern Dark Slate & Electric Blue
    BG_COLOR = RGBColor(15, 23, 42)        # Slate 900
    CARD_BG = RGBColor(30, 41, 59)         # Slate 800
    CARD_BORDER = RGBColor(51, 65, 85)     # Slate 700
    TEXT_LIGHT = RGBColor(248, 250, 252)   # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)   # Slate 400
    PRIMARY = RGBColor(56, 189, 248)       # Sky 400
    ACCENT_GREEN = RGBColor(52, 211, 153)  # Emerald 400
    ACCENT_RED = RGBColor(248, 113, 113)   # Red 400
    ACCENT_PURPLE = RGBColor(167, 139, 250)# Violet 400
    ACCENT_AMBER = RGBColor(251, 191, 36)  # Amber 400
    TAG_BG = RGBColor(30, 58, 138)         # Blue 900

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    def add_header(slide, badge_text, title_text, slide_number):
        badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.35))
        tf_b = badge_box.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = badge_text.upper()
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.font.color.rgb = PRIMARY

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(10), Inches(0.7))
        tf_t = title_box.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_LIGHT

        num_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.5), Inches(1.2), Inches(0.4))
        tf_n = num_box.text_frame
        p_n = tf_n.paragraphs[0]
        p_n.alignment = PP_ALIGN.RIGHT
        p_n.text = f"Slide {slide_number} / 5"
        p_n.font.size = Pt(12)
        p_n.font.color.rgb = TEXT_MUTED

    # ==========================================================
    # SLIDE 1: PROJECT INTRODUCTION
    # ==========================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_background(s1)

    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(0.12), Inches(4.7))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY
    bar.line.fill.background()

    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.0), Inches(11.0), Inches(0.4))
    tf = t_box.text_frame
    p = tf.paragraphs[0]
    p.text = "3RD SEMESTER MCA • PROJECT VIVA-VOCE PRESENTATION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    t_box2 = s1.shapes.add_textbox(Inches(1.2), Inches(1.4), Inches(11.0), Inches(1.3))
    tf2 = t_box2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "TMS: Complaint & Facility\nManagement System"
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_LIGHT

    t_box3 = s1.shapes.add_textbox(Inches(1.2), Inches(3.0), Inches(11.0), Inches(0.7))
    tf3 = t_box3.text_frame
    p3 = tf3.paragraphs[0]
    p3.text = "A Smart Full-Stack MERN Web Portal & React Native (Expo) Mobile App for Campus Issue Resolution"
    p3.font.size = Pt(16)
    p3.font.color.rgb = TEXT_MUTED

    # Tech Pills
    pills = ["MERN Stack", "React Native", "Expo", "JWT Auth", "REST API", "Role-Based Access"]
    pill_x = 1.2
    for pill in pills:
        pill_w = len(pill) * 0.16 + 0.55
        p_shape = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(pill_x), Inches(3.9), Inches(pill_w), Inches(0.38))
        p_shape.fill.solid()
        p_shape.fill.fore_color.rgb = TAG_BG
        p_shape.line.color.rgb = PRIMARY
        p_shape.line.width = Pt(1)
        p_tf = p_shape.text_frame
        p_p = p_tf.paragraphs[0]
        p_p.alignment = PP_ALIGN.CENTER
        p_p.text = pill
        p_p.font.size = Pt(10)
        p_p.font.bold = True
        p_p.font.color.rgb = PRIMARY
        pill_x += pill_w + 0.15

    # Candidate Card with Madan L & 1P25MC030
    pres_card = add_card(s1, Inches(1.2), Inches(4.7), Inches(10.8), Inches(1.9), CARD_BG, PRIMARY)
    pres_box = s1.shapes.add_textbox(Inches(1.5), Inches(4.85), Inches(10.2), Inches(1.6))
    pres_tf = pres_box.text_frame
    
    p = pres_tf.paragraphs[0]
    p.text = "STUDENT & PROJECT CREDENTIALS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    p_info = pres_tf.add_paragraph()
    p_info.space_before = Pt(8)
    r_name = p_info.add_run()
    r_name.text = "• Name: "
    r_name.font.bold = True
    r_name.font.color.rgb = TEXT_MUTED
    r_name_val = p_info.add_run()
    r_name_val.text = "Madan L        "
    r_name_val.font.bold = True
    r_name_val.font.size = Pt(15)
    r_name_val.font.color.rgb = TEXT_LIGHT

    r_reg = p_info.add_run()
    r_reg.text = "• Register Number: "
    r_reg.font.bold = True
    r_reg.font.color.rgb = TEXT_MUTED
    r_reg_val = p_info.add_run()
    r_reg_val.text = "1P25MC030"
    r_reg_val.font.bold = True
    r_reg_val.font.size = Pt(15)
    r_reg_val.font.color.rgb = PRIMARY

    p_guide = pres_tf.add_paragraph()
    p_guide.space_before = Pt(6)
    p_guide.text = "• Semester: 3rd Semester MCA   |   • Project Domain: Full-Stack Web & Cross-Platform Mobile"
    p_guide.font.size = Pt(12)
    p_guide.font.color.rgb = TEXT_MUTED

    # ==========================================================
    # SLIDE 2: PROBLEM & SOLUTION
    # ==========================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_background(s2)
    add_header(s2, "Motivation & Target", "Problem & Solution", 2)

    # Left: The Problem
    c_prob = add_card(s2, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.4), CARD_BG, ACCENT_RED)
    box1 = s2.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.0), Inches(5.0))
    tf1 = box1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "⚠️ The Problem (Old Manual System)"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED

    prob_points = [
        ("Lost Complaints: ", "Paper register entries easily get misplaced or forgotten."),
        ("No Live Tracking: ", "Users have no idea who is fixing their issue or when it will be done."),
        ("No Photo Proof: ", "Verbal reports fail to capture exact error codes or physical hardware damage."),
        ("Communication Delay: ", "Manual assignment of technicians takes days instead of minutes.")
    ]
    for heading, desc in prob_points:
        pt = tf1.add_paragraph()
        pt.space_before = Pt(14)
        r1 = pt.add_run()
        r1.text = "• " + heading
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = TEXT_LIGHT
        r2 = pt.add_run()
        r2.text = desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_MUTED

    # Right: The Solution
    c_sol = add_card(s2, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4), CARD_BG, ACCENT_GREEN)
    box2 = s2.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(5.0))
    tf2 = box2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "🎯 The Solution (TMS App)"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    sol_points = [
        ("100% Digital Tickets: ", "Raise complaints instantly from mobile app or web portal."),
        ("Live Progress Tracking: ", "Clear stages: Pending ➔ Assigned ➔ In-Progress ➔ Resolved."),
        ("Photo Evidence: ", "Attach photos directly using device camera to verify the problem."),
        ("Exact Location Tagging: ", "Precise hierarchy: Department ➔ Block ➔ Room.")
    ]
    for heading, desc in sol_points:
        pt = tf2.add_paragraph()
        pt.space_before = Pt(14)
        r1 = pt.add_run()
        r1.text = "✓ " + heading
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = TEXT_LIGHT
        r2 = pt.add_run()
        r2.text = desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_MUTED

    # ==========================================================
    # SLIDE 3: TECHNOLOGY USED
    # ==========================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_background(s3)
    add_header(s3, "Architecture & Frameworks", "Technology Used", 3)

    tech_cards = [
        {
            "title": "1. Frontend (Client)",
            "color": PRIMARY,
            "x": 0.8,
            "items": [
                ("React 18 (Web):", "Single-Page Web Portal for SuperAdmin management."),
                ("React Native + Expo:", "Cross-platform mobile app for iOS & Android."),
                ("Axios HTTP:", "Smooth API communication with local token caching.")
            ]
        },
        {
            "title": "2. Backend (Server)",
            "color": ACCENT_PURPLE,
            "x": 4.9,
            "items": [
                ("Node.js & Express.js:", "Fast, non-blocking asynchronous REST API."),
                ("JWT Authentication:", "Secure stateless token-based user login."),
                ("Role-Based Access (RBAC):", "Strict access for SuperAdmin, Staff, and Users.")
            ]
        },
        {
            "title": "3. Database & Security",
            "color": ACCENT_GREEN,
            "x": 9.0,
            "items": [
                ("MongoDB NoSQL:", "High-speed document database for tickets & masters."),
                ("Mongoose ODM:", "Structured data models with schema validations."),
                ("bcryptjs Hashing:", "Salted password encryption for high security.")
            ]
        }
    ]

    for tc in tech_cards:
        card = add_card(s3, Inches(tc["x"]), Inches(1.5), Inches(3.7), Inches(5.4), CARD_BG, tc["color"])
        box = s3.shapes.add_textbox(Inches(tc["x"] + 0.2), Inches(1.7), Inches(3.3), Inches(5.0))
        tf = box.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tc["title"]
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = tc["color"]

        for heading, body in tc["items"]:
            pt = tf.add_paragraph()
            pt.space_before = Pt(16)
            r1 = pt.add_run()
            r1.text = heading + " "
            r1.font.bold = True
            r1.font.size = Pt(12)
            r1.font.color.rgb = TEXT_LIGHT
            r2 = pt.add_run()
            r2.text = body
            r2.font.size = Pt(11)
            r2.font.color.rgb = TEXT_MUTED

    # ==========================================================
    # SLIDE 4: HOW TMS WORKS
    # ==========================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_background(s4)
    add_header(s4, "System Execution Flow", "How TMS Works", 4)

    steps = [
        ("Step 1", "Raise Ticket", "User opens the mobile app, selects Department ➔ Block ➔ Room, picks category, writes remarks, and snaps a photo.", PRIMARY),
        ("Step 2", "Admin Assigns", "SuperAdmin reviews incoming tickets in real-time on the Web Dashboard and assigns an available technician.", ACCENT_PURPLE),
        ("Step 3", "Resolve on Site", "Technician checks assigned ticket on Mobile App, visits the location, updates status to 'In-Progress', and fixes it.", ACCENT_AMBER),
        ("Step 4", "Close & Report", "Technician marks ticket as 'Resolved'. System saves timestamped audit trail and generates resolution analytics.", ACCENT_GREEN)
    ]

    step_w = 2.75
    s_x = 0.8
    for s_num, s_title, s_desc, s_color in steps:
        s_card = add_card(s4, Inches(s_x), Inches(1.6), Inches(step_w), Inches(5.2), CARD_BG, s_color)
        s_box = s4.shapes.add_textbox(Inches(s_x + 0.2), Inches(1.8), Inches(step_w - 0.4), Inches(4.8))
        s_tf = s_box.text_frame
        s_tf.word_wrap = True

        p1 = s_tf.paragraphs[0]
        p1.text = s_num.upper()
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = s_color

        p2 = s_tf.add_paragraph()
        p2.space_before = Pt(4)
        p2.text = s_title
        p2.font.size = Pt(17)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_LIGHT

        p3 = s_tf.add_paragraph()
        p3.space_before = Pt(16)
        p3.text = s_desc
        p3.font.size = Pt(12)
        p3.font.color.rgb = TEXT_MUTED

        s_x += step_w + 0.24

    # ==========================================================
    # SLIDE 5: BENEFITS & FUTURE SCOPE
    # ==========================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_background(s5)
    add_header(s5, "Impact & Vision", "Benefits & Future Scope", 5)

    # Left: Benefits
    c_ben = add_card(s5, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.4), CARD_BG, PRIMARY)
    b_box1 = s5.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.0), Inches(5.0))
    tf1 = b_box1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "⭐ Key Benefits of TMS"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    bens = [
        ("Zero Paperwork: ", "Saves time and eliminates misplaced physical records."),
        ("Fast Resolution: ", "Direct routing speeds up turnaround time by up to 80%."),
        ("Total Transparency: ", "Both users and management see real-time progress."),
        ("Dynamic IP Switcher: ", "Test mobile app with local backend on-the-fly without rebuilding.")
    ]
    for heading, desc in bens:
        pt = tf1.add_paragraph()
        pt.space_before = Pt(14)
        r1 = pt.add_run()
        r1.text = "• " + heading
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = TEXT_LIGHT
        r2 = pt.add_run()
        r2.text = desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_MUTED

    # Right: Future Scope
    c_fut = add_card(s5, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4), CARD_BG, ACCENT_PURPLE)
    b_box2 = s5.shapes.add_textbox(Inches(7.1), Inches(1.7), Inches(5.1), Inches(5.0))
    tf2 = b_box2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "🚀 Future Enhancements"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE

    futs = [
        ("Push Notifications: ", "Instant alerts via Firebase (FCM) when ticket status updates."),
        ("SLA Escalation Engine: ", "Auto-escalates unresolved tickets to HOD after 24 hours."),
        ("AI Classification: ", "Auto-detects problem category and assigns technician via AI."),
        ("Live Chat: ", "Direct in-app messaging between user and technician.")
    ]
    for heading, desc in futs:
        pt = tf2.add_paragraph()
        pt.space_before = Pt(14)
        r1 = pt.add_run()
        r1.text = "★ " + heading
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = TEXT_LIGHT
        r2 = pt.add_run()
        r2.text = desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_MUTED

    # Save presentation
    output_path = os.path.join(os.getcwd(), "TMS_Project_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation successfully updated at: {output_path}")

if __name__ == "__main__":
    create_presentation()
